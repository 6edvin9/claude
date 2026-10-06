"""Shipping data sheet: one row per active Shopify variant, to fill with weight and box size."""
import csv, collections, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from structure import KEEP_SHOPIFY

new = list(csv.DictReader(open('out/shopify-products-all.csv', encoding='utf-8')))
old = list(csv.DictReader(open('products_export.csv', encoding='utf-8-sig')))

def variants(rows, keep):
    out, cur = [], {}
    for r in rows:
        if r['Title']:
            cur = r
        if not keep(cur) or not r['Variant Price']:
            continue
        group = ''
        tags = [t.strip() for t in cur['Tags'].split(',')]
        sysname = cur['Type']
        opt = r['Option1 Value'] if r['Option1 Value'] not in ('', 'Default Title') else ''
        out.append([cur['Handle'], cur['Title'], sysname, opt, r['Variant SKU'], float(r['Variant Price'])])
    return out

rows = variants(new, lambda c: c.get('Status') == 'active')
rows += [[h, t, 'Patio Covers & Pergolas', o, s, p] for h, t, _, o, s, p in variants(old, lambda c: c.get('Handle') in KEEP_SHOPIFY and c.get('Status') == 'active')]

wb = Workbook()
ws = wb.active
ws.title = 'Shipping data'
F = 'Arial'
head = ['Handle (do not edit)', 'Product', 'System', 'Option', 'SKU', 'Price ($)',
        'Weight (lb)', 'Box length (in)', 'Box width (in)', 'Box height (in)', 'Boxes per order unit', 'Ships by', 'Notes',
        'Dim. weight (lb)', 'Billable weight (lb)']
ws.append(head)
fill_in = PatternFill('solid', fgColor='FFF2CC')
head_fill = PatternFill('solid', fgColor='1F2626')
grey = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='D0D0D0')
for c in ws[1]:
    c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = head_fill
    c.alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[1].height = 32
for i, r in enumerate(rows, start=2):
    ws.append(r + [None, None, None, None, 1, None, None])
    ws.cell(i, 14, '=IF(AND(H{0}>0,I{0}>0,J{0}>0),ROUND(H{0}*I{0}*J{0}/139,1),"")'.format(i))
    ws.cell(i, 15, '=IF(G{0}="","",IF(N{0}="",G{0},MAX(G{0},N{0})))'.format(i))
last = len(rows) + 1
for row in ws.iter_rows(min_row=2, max_row=last):
    for c in row:
        c.font = Font(name=F, size=10)
        c.border = Border(bottom=thin)
        if 7 <= c.column <= 13:
            c.fill = fill_in
            c.font = Font(name=F, size=10, color='0000FF')
        elif c.column in (1, 3, 5):
            c.fill = grey
    row[5].number_format = '$#,##0.00'
    for k in (6, 7, 8, 9):
        row[k].number_format = '0.0'
    row[13].number_format = '0.0'; row[14].number_format = '0.0'
dv = DataValidation(type='list', formula1='"Parcel,Freight (LTL),Local pickup / delivery"', allow_blank=True)
ws.add_data_validation(dv); dv.add('L2:L%d' % last)
for col in 'GHIJ':
    v = DataValidation(type='decimal', operator='between', formula1='0', formula2='5000', allow_blank=True,
                       error='Enter a number (pounds or inches)', showErrorMessage=True)
    ws.add_data_validation(v); v.add('%s2:%s%d' % (col, col, last))
widths = [30, 52, 18, 26, 20, 11, 11, 12, 11, 12, 11, 20, 28, 12, 12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = 'A1:O%d' % last
ws['G1'].comment = Comment('Packed weight of ONE unit as sold (the variant), in pounds, including packaging.', 'Aluglobus')
ws['H1'].comment = Comment('Outside dimensions of the shipping box in inches. If one unit ships in several boxes, enter the largest box here and the number of boxes in "Boxes per order unit".', 'Aluglobus')
ws['N1'].comment = Comment('Calculated: L x W x H / 139 (UPS/FedEx domestic divisor). Carriers bill the larger of real and dimensional weight.', 'Aluglobus')

# Guide sheet
g = wb.create_sheet('How to fill', 0)
lines = [
    ('Shipping data – how to fill', True),
    ('', False),
    ('1. Go to the "Shipping data" tab. Fill only the yellow columns (blue text): Weight, Box length / width / height, Boxes, Ships by, Notes.', False),
    ('2. One row = one thing a customer can add to the cart (each colour / configuration is its own row).', False),
    ('3. Weight = packed weight of ONE unit in pounds, including packaging. Box = outside size of the shipping box in inches.', False),
    ('4. Kits that ship in several boxes: enter the total weight, the largest box, and the number of boxes. Shopify uses one package per item, so we will set it to the total.', False),
    ('5. Ships by: Parcel (UPS/FedEx/USPS size), Freight (LTL) for long or heavy items (for example 18-24 ft profiles, gate kits on a pallet), or Local pickup / delivery.', False),
    ('6. Do not edit the Handle column – it is how the import finds the product. Grey columns are for reference.', False),
    ('7. Dim. weight and Billable weight calculate by themselves (L x W x H / 139). This is what UPS/FedEx actually bill.', False),
    ('8. Leave a row empty if you do not know yet – empty rows are skipped by the import.', False),
    ('9. Send the file back; I will turn it into a Shopify import that only updates weight and box size (nothing else changes).', False),
    ('', False),
    ('Example (how a filled row looks):', True),
]
for t, b in lines:
    g.append([t]); g.cell(g.max_row, 1).font = Font(name=F, bold=b, size=13 if t.startswith('Shipping') else 11)
g.append(['Handle', 'Product', 'Option', 'Weight (lb)', 'Box L (in)', 'Box W (in)', 'Box H (in)', 'Boxes', 'Ships by', 'Notes'])
for c in g[g.max_row]:
    c.font = Font(name=F, bold=True)
g.append(['example-gate-kit', 'Example 4x6 ft gate kit', 'Black', 68, 80, 14, 8, 2, 'Parcel', 'Frame box + slat box (example only)'])
for c in g[g.max_row]:
    c.font = Font(name=F, italic=True, color='808080')
g.append([])
g.append(['Products in the sheet: %d rows (%d products), including the 7 pergola / patio cover kits.' % (len(rows), len({r[0] for r in rows}))])
g.cell(g.max_row, 1).font = Font(name=F)
g.column_dimensions['A'].width = 24
for col in 'BCDEFGHIJ':
    g.column_dimensions[col].width = 14
g.column_dimensions['B'].width = 26
for row in g.iter_rows(min_row=1, max_row=12):
    row[0].alignment = Alignment(wrap_text=False)
wb.save('out/aluglobus-shipping-dimensions.xlsx')
print(len(rows), 'rows', len({r[0] for r in rows}), 'products')
